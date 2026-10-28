"""Convert docs/MIDTERM_REPORT.ko.md to a submission .docx (A4, Malgun Gothic; needs python-docx).
Usage: python tools/report_to_docx.py docs/MIDTERM_REPORT.ko.md out.docx [cover.json]
cover.json (optional, kept out of the repository) adds a cover page: {"학교":..., "학과":..., "학번":..., "성명":..., "지도교수":..., "제출일":...}"""
import json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
COVER = json.loads(Path(sys.argv[3]).read_text(encoding='utf-8')) if len(sys.argv) > 3 else None
FONT = '맑은 고딕'
TEXT_W = 16.0  # cm, A4 minus 2.5 cm margins

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.5))
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)

def set_font(style_or_run, size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = FONT
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, 'get_or_add_rPr') else style_or_run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(k), FONT)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color

normal = doc.styles['Normal']
set_font(normal, 10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(4)
for name, size, before, after in (('Title', 17, 0, 6), ('Heading 1', 13.5, 14, 6), ('Heading 2', 11.5, 10, 4)):
    st = doc.styles[name]
    set_font(st, size, True, RGBColor(0, 0, 0))
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
    st.paragraph_format.keep_with_next = True
    if name == 'Title':
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pbdr = st.element.get_or_add_pPr().find(qn('w:pBdr'))
        if pbdr is not None: pbdr.getparent().remove(pbdr)
for name in ('List Bullet', 'List Bullet 2'):
    set_font(doc.styles[name], 10.5)

# Page number footer
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in (('begin', None), (None, 'PAGE'), ('end', None)):
    if tag:
        el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag)
    else:
        el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = text
    run._r.append(el)
set_font(run, 9)

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def add_inline(p, text, size=None, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part: continue
        b, i, t = bold, italic, part
        if not INLINE.fullmatch(part): pass
        elif part.startswith('**'): b, t = True, part[2:-2]
        elif part.startswith('`'): t = part[1:-1]
        elif part.startswith('['): t = re.match(r'\[([^\]]+)\]', part).group(1)
        elif part.startswith('*') and len(part) > 2: i, t = True, part[1:-1]
        r = p.add_run(t); r.bold = b or None; r.italic = i or None
        set_font(r, size)

def cell_shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)

def add_table(rows):
    cols = len(rows[0])
    weights = [max(len(re.sub(r'[*`]', '', r[c])) for r in rows) + 4 for c in range(cols)]
    total = sum(weights); widths = [TEXT_W * w / total for w in weights]
    t = doc.add_table(rows=len(rows), cols=cols); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = t.cell(ri, ci); cell.width = Cm(widths[ci])
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            add_inline(p, text.strip(), size=9.5, bold=(ri == 0))
            if ri == 0: cell_shade(cell, 'E7ECF2')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_numbered(num, text, level):
    p = doc.add_paragraph()
    left = 0.7 + 0.6 * level
    p.paragraph_format.left_indent = Cm(left); p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, f'{num}. ')
    add_inline(p, text)

def add_bullet(text, level):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)

def add_cover(lines, cover):
    """Cover page from the report's own title line plus the student details; the body then starts at the summary."""
    sec.different_first_page_header_footer = True
    title = next(l for l in lines if '제목(가제)' in l).split(':', 1)[1].strip(' *')
    repo = next((l.split(':', 1)[1].strip() for l in lines if '코드 및 데이터' in l), None)
    def centered(text, size, bold=False, before=0, after=0):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text); r.bold = bold or None; set_font(r, size)
    centered('졸업논문 중간보고서', 16, True, before=80, after=36)
    centered(title, 19, True, after=6)
    centered('(가제)', 11, after=100)
    rows = [(k, cover.get(k, '')) for k in ('학교', '학과', '학번', '성명', '지도교수', '제출일') if cover.get(k)]
    t = doc.add_table(rows=len(rows), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for ri, (k, v) in enumerate(rows):
        for ci, (text, width, align) in enumerate(((k, 3.2, WD_ALIGN_PARAGRAPH.DISTRIBUTE), (v, 6.3, WD_ALIGN_PARAGRAPH.LEFT))):
            cell = t.cell(ri, ci); cell.width = Cm(width)
            p = cell.paragraphs[0]; p.alignment = align; p.paragraph_format.space_after = Pt(6)
            r = p.add_run(text); set_font(r, 12, ci == 0)
    if repo:
        centered('코드 및 데이터: ' + repo, 9.5, before=60)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

lines = SRC.read_text(encoding='utf-8').splitlines()
if COVER:
    add_cover(lines, COVER)
    lines = lines[next(k for k, l in enumerate(lines) if l.startswith('## ')):]
i = 0; table = []; in_refs = False
while i < len(lines):
    line = lines[i]; stripped = line.strip()
    if table and not stripped.startswith('|'):
        add_table(table); table = []
    if not stripped:
        i += 1; continue
    if stripped.startswith('|'):
        cells = [c for c in stripped.strip('|').split('|')]
        if not all(re.fullmatch(r'\s*:?-+:?\s*', c) for c in cells): table.append(cells)
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '):
        in_refs = line[3:].strip() == '참고 문헌'
        doc.add_paragraph(line[3:], style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], style='Heading 2')
    elif stripped.startswith('!['):
        path = (SRC.parent / re.search(r'\]\(([^)]+)\)', stripped).group(1)).resolve()
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        width = 13.5 if 'method' in path.name else 11.0 if 'grids' in path.name else TEXT_W
        p.add_run().add_picture(str(path), width=Cm(width))
    elif stripped.startswith('*그림'):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
        add_inline(p, stripped.strip('*'), size=9.5)
    elif m := re.match(r'^(\s*)-\s+(.*)', line):
        add_bullet(m.group(2), len(m.group(1)) // 2)
    elif m := re.match(r'^(\s*)(\d+)\.\s+(.*)', line):
        if in_refs:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.2
            add_inline(p, f'[{m.group(2)}] ', size=9.5); add_inline(p, m.group(3), size=9.5)
        else:
            add_numbered(m.group(2), m.group(3), len(m.group(1)) // 3)
    else:
        p = doc.add_paragraph(); add_inline(p, stripped)
    i += 1
if table: add_table(table)
doc.core_properties.title = '졸업논문 중간보고서'
doc.save(OUT)
print('saved', OUT)
